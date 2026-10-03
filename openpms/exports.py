"""Exports: tidy files for spreadsheets, BI tools and AI.

Every export (a download, a scheduled snapshot or a data link) is one of the
DATASETS below. Each dataset is a set of tidy tables: one row per record,
snake_case headers, ISO dates, numbers as numbers, blanks as blanks, and
lookup columns (names, labels) filled in so each file stands on its own.

Every export also carries a data dictionary and a README written for AI
tools, and the quarter pack carries a prompt for drafting the quarterly
report. Nothing here depends on any one AI product: the files work with any
of them, or with none.

Never exported: individual wellbeing, reviewer comments, the audit log,
drafts and unapproved values, and line manager links.
"""
import calendar
import csv
import io
import os
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from . import __version__, db, periods as P
from .schema import LISTS_BY_NAME

TYPE_MAP = {"text": "text", "note": "text", "choice": "text", "number": "number",
            "integer": "number", "bool": "bool", "date": "date", "datetime": "datetime"}
FORMATS = {"xlsx": "Excel workbook", "csv": "CSV files (zip)", "both": "CSV and Excel (zip)"}
FOLDER_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,59}$")


# ---------------------------------------------------------------------------
# Column and table definitions
# ---------------------------------------------------------------------------
def _cols(list_name, include=None, exclude=()):
    lst = LISTS_BY_NAME[list_name]
    out = [dict(name=c.name, source=c.name, type=TYPE_MAP[c.type], description=c.description)
           for c in lst.all_columns if (not include or c.name in include) and c.name not in exclude]
    if include:
        order = {n: i for i, n in enumerate(include)}
        out.sort(key=lambda c: order[c["name"]])
    return out


def _lookup(name, from_field, tables, description):
    """A column filled from another table: [(table, field)], tried in order,
    matched on that table's key."""
    return dict(name=name, lookup_from=from_field, type="text", description=description, lookup=tables)


PERSON = [("people", "display_name")]
TEAM = [("org_units", "unit_name")]
CONTRIB = [("groups", "group_name"), ("measures", "measure_name")]
LOOKUP_TABLES = {"people": "display_name", "org_units": "unit_name", "groups": "group_name", "measures": "measure_name"}

WEEKLY_UPDATE_COLS = (
    _cols("weekly_updates", ["update_key", "week_start", "email"])
    + [_lookup("person_name", "email", PERSON, "Name of the person."),
       *_cols("weekly_updates", ["org_unit_key"]),
       _lookup("team_name", "org_unit_key", TEAM, "Team name.")]
    + _cols("weekly_updates", ["successes", "communication", "workload", "submitted_at"])
)
TASK_COLS = (
    _cols("tasks", ["task_code", "task_name", "priority", "status", "org_unit_key"])
    + [_lookup("team_name", "org_unit_key", TEAM, "Team name.")]
    + _cols("tasks", ["contributes_to_type", "contributes_to_key"])
    + [_lookup("contributes_to_name", "contributes_to_key", CONTRIB, "Name of the group or measure the task supports.")]
    + _cols("tasks", ["contributes_to_note", "raised_by"])
    + [_lookup("raised_by_name", "raised_by", PERSON, "Who raised it.")]
    + _cols("tasks", ["date_raised", "assigned_to"])
    + [_lookup("assigned_to_name", "assigned_to", PERSON, "Who is doing it.")]
    + _cols("tasks", ["completed_by"])
    + [_lookup("completed_by_name", "completed_by", PERSON, "Who completed it.")]
    + _cols("tasks", ["completed_date", "notes"])
)
PROBLEM_COLS = (
    _cols("problems", ["problem_code", "problem_title", "problem_statement", "status", "impact", "urgency", "org_unit_key"])
    + [_lookup("team_name", "org_unit_key", TEAM, "Team name.")]
    + _cols("problems", ["raised_by"])
    + [_lookup("raised_by_name", "raised_by", PERSON, "Who raised it.")]
    + _cols("problems", ["raised_on", "problem_owner"])
    + [_lookup("problem_owner_name", "problem_owner", PERSON, "Who owns resolving it.")]
    + _cols("problems", ["target_resolution_date", "closed_date", "resolution", "contributes_to_type", "contributes_to_key"])
    + [_lookup("contributes_to_name", "contributes_to_key", CONTRIB, "Name of the group or measure affected.")]
    + _cols("problems", ["notes"])
)

ABOUT = {
    "measures": "One row per measure: definition, unit, polarity, frequency and owning team.",
    "periods": "One row per reporting period: dates, financial year and quarter.",
    "measure_roles": "Who owns, updates and approves each measure.",
    "people": "People and their teams.",
    "org_units": "Teams, with their parent workstream or team.",
    "groups": "Themes, programmes and objectives that measures and work link to.",
    "measure_links": "Which measures belong to which themes, programmes and objectives.",
    "reference_values": "Targets, tolerances, baselines and capacities by measure and period.",
    "submissions": "One row per measure per period, with its current status.",
    "approved_values": "Every approved version of every value, including earlier approved versions that were later corrected.",
    "values_flat": "One row per approved measure per period with every lookup filled in. Start here.",
    "quarter_values": "Approved values for periods in the quarter, including the quarter itself.",
    "latest_values": "The latest approved value of every measure, which covers annual, academic-year and term measures.",
    "weekly_updates": "Successes, communication and workload from weekly updates. Wellbeing is never included.",
    "tasks": "Tasks, one row each for their whole life.",
    "problems": "Problems, one row each for their whole life.",
}


def _table(name, list_name, columns, where=None, main=False):
    return dict(name=name, list=list_name, columns=columns, where=where, main=main)


def _in_quarter(r, ctx, col):
    q, d = ctx["quarter"], r[col]
    return d is not None and q["start_date"] <= d <= q["end_date"]


DATASETS = {
    "measures": dict(
        title="Measure values",
        description="One row per approved measure per period, with measure, owner, period, targets and RAG "
                    "filled in. Works on its own with no joins.",
        tables=[_table("values_flat", "rpt_values", _cols("rpt_values"),
                       lambda r, ctx: not ctx["financial_year"] or r["financial_year"] == ctx["financial_year"],
                       main=True)],
    ),
    "full_model": dict(
        title="Full data model",
        description="Every table in the model, for BI tools, spreadsheets or AI. Approved values only; "
                    "review comments, wellbeing and the audit log are never included.",
        tables=[
            _table("measures", "measures", _cols("measures")),
            _table("periods", "periods", _cols("periods")),
            _table("measure_roles", "measure_roles", _cols("measure_roles")),
            _table("people", "people", _cols("people", exclude=("line_manager_email",))),
            _table("org_units", "org_units", _cols("org_units")),
            _table("groups", "groups", _cols("groups")),
            _table("measure_links", "measure_links", _cols("measure_links")),
            _table("reference_values", "reference_values", _cols("reference_values")),
            _table("submissions", "submissions", _cols("submissions")),
            _table("approved_values", "submission_versions",
                   [dict(c, description="approved for the version in use, superseded for one replaced by a later correction.")
                    if c["name"] == "version_status" else c
                    for c in _cols("submission_versions", exclude=("entered_by", "entered_date"))],
                   lambda r, ctx: r["version_status"] in ("approved", "superseded")),
            _table("values_flat", "rpt_values", _cols("rpt_values"), main=True),
            _table("weekly_updates", "weekly_updates", WEEKLY_UPDATE_COLS),
            _table("tasks", "tasks", TASK_COLS),
            _table("problems", "problems", PROBLEM_COLS),
        ],
    ),
    "weekly_work": dict(
        title="Weekly work",
        description="Weekly updates (successes, communication, workload), tasks and problems, with names filled in. "
                    "Wellbeing is never included.",
        tables=[
            _table("weekly_updates", "weekly_updates", WEEKLY_UPDATE_COLS, main=True),
            _table("tasks", "tasks", TASK_COLS, main=True),
            _table("problems", "problems", PROBLEM_COLS, main=True),
        ],
    ),
    "quarter_pack": dict(
        title="Quarterly report pack",
        description="Everything needed to write the quarterly report for one financial quarter: the quarter's "
                    "approved values, the latest value of every measure, work completed, open work, problems "
                    "and successes. Comes with a prompt for drafting the report with AI, and a checklist for "
                    "writing it without.",
        tables=[
            _table("quarter_values", "rpt_values", _cols("rpt_values"),
                   lambda r, ctx: r["financial_quarter"] == ctx["quarter"]["financial_quarter"]),
            _table("latest_values", "rpt_values", _cols("rpt_values"), lambda r, ctx: r["is_latest"]),
            _table("tasks", "tasks", TASK_COLS,
                   lambda r, ctx: _in_quarter(r, ctx, "completed_date") or r["status"] == "open"),
            _table("problems", "problems", PROBLEM_COLS,
                   lambda r, ctx: _in_quarter(r, ctx, "raised_on") or _in_quarter(r, ctx, "closed_date")
                   or r["status"] == "open"),
            _table("weekly_updates", "weekly_updates", WEEKLY_UPDATE_COLS,
                   lambda r, ctx: _in_quarter(r, ctx, "week_start")),
        ],
    ),
}
DATA_LINK_DATASETS = ("measures", "full_model", "weekly_work")

DICTIONARY_HEADERS = ["table_name", "column_name", "data_type", "description"]
DICTIONARY_COLUMNS = [dict(name=h, type="text") for h in DICTIONARY_HEADERS]


def dictionary_rows(dataset):
    return [[t["name"], c["name"], c["type"], c["description"]]
            for t in DATASETS[dataset]["tables"] for c in t["columns"]]


# ---------------------------------------------------------------------------
# Building an export
# ---------------------------------------------------------------------------
@dataclass
class Export:
    dataset: str
    tables: list  # [(table definition, [row values in column order])]
    created_at: datetime
    filter_label: str
    org_name: str
    quarter_label: str = None
    files: dict = field(default_factory=dict)  # text files: name -> text

    @property
    def title(self):
        return DATASETS[self.dataset]["title"]

    @property
    def base_name(self):
        return f"{self.dataset}_{self.created_at.strftime('%Y-%m-%d')}"

    def table(self, name):
        return next(rows for t, rows in self.tables if t["name"] == name)

    def row_count(self):
        return sum(len(rows) for _, rows in self.tables)


def _setting(repo, key, default):
    row = repo.get("settings", key)
    return row["setting_value"] if row and row["setting_value"] not in (None, "") else default


def fy_start_month(repo):
    v = _setting(repo, "fy_start_month", "4")
    return int(v) if str(v).isdigit() and 1 <= int(v) <= 12 else 4


def quarters(repo, today=None):
    """Financial quarters that have started, newest first."""
    today = today or date.today()
    return repo.find("periods", "period_type = 'quarterly' AND start_date <= ?", (today.isoformat(),),
                     order="start_date DESC")


def last_completed_quarter(repo, today=None):
    today = today or date.today()
    rows = repo.find("periods", "period_type = 'quarterly' AND end_date < ?", (today.isoformat(),),
                     order="end_date DESC", limit=1)
    return rows[0] if rows else None


def financial_years(repo):
    rows = repo.conn.execute("SELECT DISTINCT financial_year FROM rpt_values WHERE org_id = ? AND financial_year IS NOT NULL"
                             " ORDER BY financial_year DESC", (repo.org_id,)).fetchall()
    return [r[0] for r in rows]


def _lookup_maps(repo):
    out = {}
    for table, fld in LOOKUP_TABLES.items():
        key = LISTS_BY_NAME[table].key
        out[table] = {r[key]: r[fld] for r in repo.find(table)}
    return out


def _value(col, row, maps):
    if col.get("lookup"):
        key = row.get(col["lookup_from"])
        if key in (None, ""):
            return None
        for table, _ in col["lookup"]:
            name = maps[table].get(key)
            if name not in (None, ""):
                return name
        return None
    return row.get(col["source"])


def build(repo, dataset, *, quarter=None, financial_year=None, now=None, only=None):
    """Collect one export. quarter: a quarterly periods row (quarter_pack
    only; defaults to the last completed quarter). financial_year: e.g.
    '2026-27' (measures only). only: table names to read (default all)."""
    if dataset not in DATASETS:
        raise ValueError(f"Unknown dataset {dataset}")
    now = now or db.utcnow()
    if dataset == "quarter_pack" and quarter is None:
        quarter = last_completed_quarter(repo, now.date())
        if quarter is None:
            raise ValueError("There are no completed quarters yet.")
    ctx = dict(quarter=quarter, financial_year=financial_year if dataset == "measures" else None)
    maps = _lookup_maps(repo)
    cache, tables = {}, []
    for t in DATASETS[dataset]["tables"]:
        if only is not None and t["name"] not in only:
            continue
        if t["list"] not in cache:
            cache[t["list"]] = repo.find(t["list"])
        rows = [r for r in cache[t["list"]] if t["where"] is None or t["where"](r, ctx)]
        tables.append((t, [[_value(c, r, maps) for c in t["columns"]] for r in rows]))
    if dataset == "quarter_pack":
        label = f"financial quarter {quarter['financial_quarter']} ({quarter['start_date']} to {quarter['end_date']})"
    elif ctx["financial_year"]:
        label = f"financial year {ctx['financial_year']}"
    else:
        label = "everything"
    exp = Export(dataset, tables, now, label, _setting(repo, "organisation_name", "Open PMS"),
                 quarter["financial_quarter"] if dataset == "quarter_pack" else None)
    exp.files["README_for_AI.md"] = readme_for(dataset, fy_start_month(repo))
    if dataset == "quarter_pack":
        exp.files["prompt_quarterly_report.md"] = QUARTER_PROMPT.replace("{quarter}", quarter["financial_quarter"])
    exp.files["export_info.txt"] = _info(exp)
    return exp


def _info(exp):
    lines = [f"dataset: {exp.dataset}", f"title: {exp.title}", f"organisation: {exp.org_name}",
             f"created_at: {db.iso(exp.created_at)}", f"filter: {exp.filter_label}", f"open_pms_version: {__version__}",
             "rows:"]
    lines += [f"  {t['name']}: {len(rows)}" for t, rows in exp.tables]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Writing files
# ---------------------------------------------------------------------------
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def csv_text(v, typ):
    """One value as CSV text. Text that a spreadsheet would run as a formula
    gets a leading apostrophe, so opening an export can never run anything."""
    if v is None or v == "":
        return ""
    if typ == "bool":
        return "true" if v else "false"
    if isinstance(v, datetime):
        return db.iso(v)
    if isinstance(v, date):
        return v.isoformat()
    if typ == "number":
        n = float(v)
        return str(int(n)) if n.is_integer() else repr(n)
    s = str(v)
    return "'" + s if s.startswith(FORMULA_START) else s


def to_csv(columns, rows):
    """UTF-8 with a byte order mark (so Excel reads accents and £ correctly),
    CRLF line endings and a single header row."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([c["name"] for c in columns])
    types = [c["type"] for c in columns]
    for r in rows:
        w.writerow([csv_text(v, t) for v, t in zip(r, types)])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def csv_files(exp):
    """Every table's CSV plus the data dictionary."""
    out = {f"{t['name']}.csv": to_csv(t["columns"], rows) for t, rows in exp.tables}
    out["data_dictionary.csv"] = to_csv(DICTIONARY_COLUMNS, dictionary_rows(exp.dataset))
    return out


def to_xlsx(exp):
    """One sheet per table, then data_dictionary and readme. Real dates and
    numbers, a frozen header row and filters. Text is always stored as text,
    never as a formula."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    def sheet(name, columns, rows):
        ws = wb.create_sheet(name[:31])
        ws.freeze_panes = "A2"
        for i, c in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(48, len(c["name"]) + 4))
        header = []
        for c in columns:
            cell = WriteOnlyCell(ws, value=c["name"])
            cell.font = bold
            header.append(cell)
        ws.append(header)
        for r in rows:
            out = []
            for v, c in zip(r, columns):
                t = c["type"]
                if v is None or v == "":
                    out.append(None)
                    continue
                if t == "datetime" and isinstance(v, datetime):
                    cell = WriteOnlyCell(ws, value=v.astimezone(timezone.utc).replace(tzinfo=None))
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif t == "date" and isinstance(v, date):
                    cell = WriteOnlyCell(ws, value=v)
                    cell.number_format = "yyyy-mm-dd"
                elif t == "number":
                    cell = WriteOnlyCell(ws, value=v)
                elif t == "bool":
                    cell = WriteOnlyCell(ws, value=bool(v))
                else:
                    cell = WriteOnlyCell(ws, value=str(v))
                    cell.data_type = "s"
                out.append(cell)
            ws.append(out)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    for t, rows in exp.tables:
        sheet(t["name"], t["columns"], rows)
    sheet("data_dictionary", DICTIONARY_COLUMNS, dictionary_rows(exp.dataset))
    text = exp.files["README_for_AI.md"]
    if "prompt_quarterly_report.md" in exp.files:
        text += "\n\n" + exp.files["prompt_quarterly_report.md"]
    text += "\n\n" + exp.files["export_info.txt"]
    sheet("readme", [dict(name="text", type="text")], [[line] for line in text.splitlines()])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def files_for(exp, fmt):
    """{file name: bytes} for a format: csv, xlsx or both."""
    out = {}
    if fmt in ("csv", "both"):
        out.update(csv_files(exp))
    if fmt in ("xlsx", "both"):
        out[f"{exp.dataset}.xlsx"] = to_xlsx(exp)
    for name, text in exp.files.items():
        out[name] = text.encode("utf-8")
    return out


def download_name(exp, fmt):
    return f"{exp.base_name}.{'xlsx' if fmt == 'xlsx' else 'zip'}"


def download(exp, fmt):
    """(file name, mimetype, bytes). Excel on its own is a single workbook;
    anything with CSV files is a zip with one folder inside."""
    if fmt == "xlsx":
        return (download_name(exp, fmt), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                to_xlsx(exp))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files_for(exp, fmt).items():
            z.writestr(f"{exp.base_name}/{name}", data)
    return download_name(exp, fmt), "application/zip", buf.getvalue()


def _ref(now):
    return f"REQ-{now.strftime('%Y%m%d-%H%M')}-{os.urandom(3).hex()}"


def log_download(repo, exp, fmt, who):
    now = db.utcnow()
    repo.insert("export_requests", dict(
        request_ref=_ref(now), dataset=exp.dataset,
        quarter_label=exp.quarter_label,
        filter_label=exp.filter_label, folder_path="download", requested_by=who, requested_at=now, status="done",
        output_url=download_name(exp, fmt), message=f"{exp.row_count():,} rows, {FORMATS[fmt]}.",
        completed_at=now), who, now, audit=False)


# ---------------------------------------------------------------------------
# Scheduled snapshots
# ---------------------------------------------------------------------------
def snapshot_root(base_dir, org_slug):
    """Single-organisation installs write straight into the snapshot folder."""
    return Path(base_dir) if org_slug == "default" else Path(base_dir) / org_slug


def write_folder(path, files):
    """Replace a folder's contents in one step, so nobody reading it (or
    syncing it) ever sees half an export."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix=f".{path.name}-new-", dir=path.parent))
    for name, data in files.items():
        (tmp / name).write_bytes(data)
    old = None
    if path.exists():
        old = Path(tempfile.mkdtemp(prefix=f".{path.name}-old-", dir=path.parent))
        old.rmdir()
        path.rename(old)
    tmp.rename(path)
    if old:
        shutil.rmtree(old, ignore_errors=True)


def last_due(job, today, fy_month=4):
    """The most recent date on or before today when the job was due."""
    f, day = job["frequency"], job["run_day"]
    if f == "daily":
        return today
    if f == "weekly":
        wd = day if day and 1 <= day <= 7 else 1
        return today - timedelta(days=(today.isoweekday() - wd) % 7)
    step = 1 if f == "monthly" else 3
    y, m = today.year, today.month
    if f == "quarterly":  # quarters start in the financial year's quarter months
        y, m = P._add_months(y, m, -((m - fy_month) % 3))
    for _ in range(2):
        d = date(y, m, min(day or 1, calendar.monthrange(y, m)[1]))
        if d <= today:
            return d
        y, m = P._add_months(y, m, -step)
    return None


def job_due(job, today, fy_month=4):
    """Due if it hasn't run since its most recent run day, so a run day when
    nobody used the app is caught up the next day."""
    if not job["active"]:
        return False
    due = last_due(job, today, fy_month)
    return due is not None and (job["last_run_at"] is None or job["last_run_at"].date() < due)


def run_job(repo, job, base_dir, org_slug, now=None, who="system"):
    """Write one job's files to <folder>/latest and, if kept, <folder>/<date>.
    Never raises: the outcome is recorded on the job and in export_requests."""
    now = now or db.utcnow()
    status, where, quarter_label = "done", None, None
    try:
        if not FOLDER_NAME.match(job["folder_path"] or ""):
            raise ValueError("the folder name can only use letters, numbers, - and _")
        exp = build(repo, job["dataset"], now=now)
        quarter_label = exp.quarter_label
        files = files_for(exp, job["format"])
        root = snapshot_root(base_dir, org_slug) / job["folder_path"]
        write_folder(root / "latest", files)
        if job["keep_history"]:
            write_folder(root / now.date().isoformat(), files)
        where = f"{job['folder_path']}/latest"
        message = f"{len(exp.tables)} tables, {exp.row_count():,} rows. Filter: {exp.filter_label}."
    except Exception as exc:  # recorded so admins can see it on the job
        status, message = "failed", str(exc) or type(exc).__name__
    with repo.transaction():
        repo.update("export_jobs", job["job_code"],
                    dict(last_run_at=now, last_run_status=f"{status}: {message}"[:250]), who, now, audit=False)
        repo.insert("export_requests", dict(
            request_ref=_ref(now), dataset=job["dataset"], job_code=job["job_code"], quarter_label=quarter_label, folder_path=job["folder_path"] or "-",
            requested_at=now, status=status, output_url=where, message=message, completed_at=db.utcnow()),
            who, now, audit=False)
    return status


def run_due_jobs(repo, base_dir, org_slug, today=None, now=None):
    """Run every active job that's due. Returns {job_code: status}."""
    now = now or db.utcnow()
    today = today or now.date()
    fy = fy_start_month(repo)
    return {job["job_code"]: run_job(repo, job, base_dir, org_slug, now)
            for job in repo.find("export_jobs", "active = 1") if job_due(job, today, fy)}


# ---------------------------------------------------------------------------
# Text files written with every export
# ---------------------------------------------------------------------------
def year_rule(fy_month):
    if fy_month == 1:
        return ("- The financial year is the calendar year, January to December. `financial_year` looks like `2026` "
                "and `financial_quarter` like `2026 Q2` (Q1 is January to March).")
    first, last = calendar.month_name[fy_month], calendar.month_name[(fy_month - 2) % 12 + 1]
    q1_end = calendar.month_name[(fy_month + 1) % 12 + 1]
    return (f"- The financial year runs {first} to {last}. `financial_year` looks like `2026-27` and "
            f"`financial_quarter` like `2026-27 Q2` (Q1 is {first} to {q1_end}).")


def reading_rules(fy_month):
    return f"""## How to read this data

- Every file is tidy: one row per record, one column per field, a single header row, no merged cells and no totals.
- Column names are snake_case and mean the same thing in every file. `data_dictionary.csv` (or the `data_dictionary` sheet) describes every column.
- Dates are `YYYY-MM-DD`. Date-times are UTC in ISO 8601, for example `2026-09-30T06:00:00Z`.
- Numbers are plain numbers. A blank cell means there is no value; it never means zero.
- Yes or no columns hold `true` or `false`.
- In CSV files, text that starts with `=`, `+`, `-` or `@` has an apostrophe added at the front, so spreadsheets don't treat it as a formula. Ignore that apostrophe.
- `percent` values are 0 to 100, so 58 means 58%. `gbp` is pounds. `yes_no` is 1 (yes) or 0 (no). `text` measures use `value_text`.
- Only **approved** values are included. Drafts, returned values and reviewer comments are not.
- `measure_code` (for example PM-0007) is the unique key for a measure. `source_ref` is the reference used in the source document and can repeat.
{year_rule(fy_month)}
- `rag_status`: `green` is on or better than target; `amber` is short of target but within tolerance; `red` is off track; `no_target` means no target was set; `no_data` means no value was provided; `not_applicable` means RAG isn't used for that measure.
- `polarity` says whether higher or lower is better. `aggregation_method` says how to combine values into a longer period: sum, average, latest, max or min. `do_not_combine` means the values must not be added up or averaged.
- `is_latest` is true on the most recent approved period for each measure.
- Individual wellbeing, reviewer comments and the audit log are never included in exports.
"""


JOINS = """## How the tables join

| From | Column | To |
|---|---|---|
| any table | `measure_code` | `measures.measure_code` |
| any table | `period_key` | `periods.period_key` |
| measures | `parent_measure_code` | `measures.measure_code` (parent summary measure) |
| submissions, approved_values | `submission_key` | `submissions.submission_key` |
| measure_links | `group_key` | `groups.group_key` |
| measure_roles, weekly_updates | `email` | `people.email` |
| most tables | `org_unit_key` | `org_units.org_unit_key` |
| tasks, problems | `contributes_to_key` | `groups.group_key` or `measures.measure_code` (see `contributes_to_type`) |

`values_flat` already has every lookup filled in, so for most questions it's the only table you need.
"""


def readme_for(dataset, fy_month=4):
    d = DATASETS[dataset]
    lines = [f"# {d['title']}", "", d["description"], "",
             "`export_info.txt` (or the end of the `readme` sheet) says when the export ran, what filter was used "
             "and how many rows each table has.", "", "## Files", ""]
    for t in d["tables"]:
        lines.append(f"- `{t['name']}.csv` (sheet `{t['name']}` in the Excel file): {ABOUT[t['name']]}")
    lines += ["- `data_dictionary.csv` (sheet `data_dictionary`): every column in this export, with its type and meaning.",
              "- `README_for_AI.md` (sheet `readme`): this file."]
    if dataset == "quarter_pack":
        lines.append("- `prompt_quarterly_report.md`: a prompt for drafting the quarterly report from these files, "
                     "and a checklist for writing it without AI.")
    lines += ["", reading_rules(fy_month)]
    if dataset == "full_model":
        lines.append(JOINS)
    if dataset == "quarter_pack":
        lines += ["## What's in each table", "",
                  "- `quarter_values`: approved values for periods in the quarter, including the quarter itself.",
                  "- `latest_values`: the latest approved value of every measure, which covers annual, academic-year and term measures.",
                  "- `tasks`: tasks completed in the quarter, plus everything still open.",
                  "- `problems`: problems raised or closed in the quarter, plus everything still open.",
                  "- `weekly_updates`: successes, communication and workload from every weekly update in the quarter.", ""]
    return "\n".join(lines)


QUARTER_PROMPT = """# Drafting the quarterly performance report for {quarter}

## With AI

Use any AI tool your organisation allows, for example Microsoft Copilot, ChatGPT, Gemini, Claude or a model you run yourselves. Nothing here depends on one product, so if a tool changes or goes away, use another. Attach the files in this folder (or the Excel file), then paste everything below the line.

---

You are helping write the quarterly performance report for {quarter}. Use **only** the attached files. Read `README_for_AI.md` and `data_dictionary.csv` (or the `readme` and `data_dictionary` sheets) first, so you understand every column.

Write the report in plain UK English, for senior leaders who have two minutes. Use these sections:

1. **Summary** (five bullet points at most): overall position, the biggest improvement, the biggest concern, and what needs a decision.
2. **Measures**, grouped by `category`:
   - a table with measure (`source_ref` and `measure_name`), latest value with its unit, target, RAG and the direction of travel from the previous period;
   - for every red or amber measure, one or two sentences from its `narrative` explaining why and what is being done.
3. **Delivery**: the main work completed this quarter (`tasks` with `status` complete), grouped by what it `contributes_to_name`, and notable successes from `weekly_updates`.
4. **Problems and risks**: open problems with high impact or high urgency, who owns them and their target date, then problems resolved this quarter.
5. **Data notes**: measures with `no_data`, values marked `provisional` or `estimated` in `data_quality`, and measures with no target.

Rules:
- Never invent or estimate numbers. If something isn't in the files, say "not available".
- Quote measure codes (for example PM-0007) so every statement can be checked.
- Percentages are stored as 0 to 100: show 58 as 58%.
- Use `polarity` when you describe change: for lower-is-better measures, a fall is an improvement.
- Use `aggregation_method` if you combine monthly or weekly values into a quarter figure, and say that you did.
- Don't name individuals when you describe problems or workload; use team names.
- Finish with a short list of questions the report owner should check before publishing.

---

## Without AI

Open the Excel file and work through it in this order. Each step matches a section of the report.

1. **Measures**: on `latest_values`, filter `rag_status` to red and amber and sort by `category`. Copy `source_ref`, `measure_name`, `value_number` (or `value_text`), `target_value` and `rag_status` into the report table, and use each `narrative` for the explanation.
2. **Delivery**: on `tasks`, filter `status` to complete and sort by `contributes_to_name`. Pick out notable successes from the `successes` column of `weekly_updates`.
3. **Problems and risks**: on `problems`, filter `status` to open and `impact` or `urgency` to high.
4. **Data notes**: on `latest_values`, filter `rag_status` to no_data or no_target, and `data_quality` to provisional or estimated.
5. **Summary**: write it last, from the four sections above.
"""
