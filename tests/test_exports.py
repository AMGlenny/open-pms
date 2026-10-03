"""Exports: the files themselves, downloads, scheduled snapshots and data links."""
import csv
import io
import re
import unittest
import zipfile
from datetime import date, datetime, timezone

from openpyxl import load_workbook

from openpms import db, exports

from .helpers import ADMIN, AppCase, STANDARD, VIEWER

NOW = datetime(2026, 10, 3, 9, 0, tzinfo=timezone.utc)


def read_csv(data):
    text = data.decode("utf-8")
    assert text.startswith("﻿"), "byte order mark so Excel reads UTF-8"
    return list(csv.reader(io.StringIO(text[1:])))


class Values(unittest.TestCase):
    def test_csv_text(self):
        self.assertEqual(exports.csv_text(None, "text"), "")
        self.assertEqual(exports.csv_text(58.0, "number"), "58")
        self.assertEqual(exports.csv_text(12.5, "number"), "12.5")
        self.assertEqual(exports.csv_text(0, "number"), "0", "zero is a value, not a blank")
        self.assertEqual(exports.csv_text(True, "bool"), "true")
        self.assertEqual(exports.csv_text(False, "bool"), "false")
        self.assertEqual(exports.csv_text(date(2026, 9, 30), "date"), "2026-09-30")
        self.assertEqual(exports.csv_text(NOW, "datetime"), "2026-10-03T09:00:00Z")

    def test_formulas_never_run_in_spreadsheets(self):
        for s in ("=HYPERLINK(\"http://x\")", "+1", "-5 below target", "@SUM(A1)"):
            self.assertEqual(exports.csv_text(s, "text"), "'" + s)
        self.assertEqual(exports.csv_text("Fine text", "text"), "Fine text")
        self.assertEqual(exports.csv_text(-5, "number"), "-5", "negative numbers stay numbers")

    def test_csv_quoting_round_trips(self):
        cols = [dict(name="a", type="text"), dict(name="b", type="number")]
        rows = [["Comma, \"quotes\" and\na new line", 3]]
        self.assertEqual(read_csv(exports.to_csv(cols, rows)), [["a", "b"], ["Comma, \"quotes\" and\na new line", "3"]])

    def test_year_rule_follows_settings(self):
        self.assertIn("runs April to March", exports.year_rule(4))
        self.assertIn("Q1 is April to June", exports.year_rule(4))
        self.assertIn("runs July to June", exports.year_rule(7))
        self.assertIn("calendar year", exports.year_rule(1))


class Schedule(unittest.TestCase):
    def job(self, frequency, run_day=None, last=None):
        return dict(active=True, frequency=frequency, run_day=run_day, last_run_at=last)

    def test_last_due(self):
        sat = date(2026, 10, 3)
        self.assertEqual(exports.last_due(self.job("daily"), sat), sat)
        self.assertEqual(exports.last_due(self.job("weekly", 1), sat), date(2026, 9, 28))
        self.assertEqual(exports.last_due(self.job("monthly", 10), sat), date(2026, 9, 10))
        self.assertEqual(exports.last_due(self.job("monthly", 31), date(2026, 3, 1)), date(2026, 2, 28), "short months")
        self.assertEqual(exports.last_due(self.job("quarterly", 10), sat), date(2026, 7, 10))
        self.assertEqual(exports.last_due(self.job("quarterly", 10), date(2026, 10, 10)), date(2026, 10, 10))
        self.assertEqual(exports.last_due(self.job("quarterly", 10), date(2026, 10, 10), fy_month=1), date(2026, 10, 10))
        self.assertEqual(exports.last_due(self.job("quarterly", 10), date(2026, 9, 10), fy_month=6), date(2026, 9, 10))

    def test_missed_run_day_is_caught_up(self):
        ran_last_week = datetime(2026, 9, 21, 6, tzinfo=timezone.utc)
        self.assertTrue(exports.job_due(self.job("weekly", 1, ran_last_week), date(2026, 9, 29)),
                        "nobody used the app on Monday, so it runs on Tuesday")
        ran_monday = datetime(2026, 9, 28, 6, tzinfo=timezone.utc)
        self.assertFalse(exports.job_due(self.job("weekly", 1, ran_monday), date(2026, 9, 29)))
        self.assertFalse(exports.job_due(dict(self.job("daily"), active=False), date(2026, 9, 29)))


class Build(AppCase, unittest.TestCase):
    def test_measures_matches_reporting_table(self):
        repo = self.repo()
        exp = exports.build(repo, "measures", now=NOW)
        rows = exp.table("values_flat")
        self.assertEqual(len(rows), repo.count("rpt_values"))
        headers = read_csv(exports.csv_files(exp)["values_flat.csv"])[0]
        self.assertEqual(headers[:3], ["submission_key", "measure_code", "measure_name"])
        self.assertTrue(all(re.fullmatch(r"[a-z][a-z0-9_]*", h) for h in headers), "snake_case headers")
        fy = exports.build(repo, "measures", financial_year="2026-27", now=NOW)
        self.assertTrue(0 < len(fy.table("values_flat")) < len(rows))
        self.assertIn("financial year 2026-27", fy.files["export_info.txt"])

    def test_private_data_never_exported(self):
        exp = exports.build(self.repo(), "full_model", now=NOW)
        names = {t["name"] for t, _ in exp.tables}
        self.assertFalse(names & {"wellbeing_checkins", "review_comments", "audit_log", "export_requests"})
        everything = b"".join(exports.files_for(exp, "both").values())
        self.assertNotIn(b"line_manager_email", everything)
        self.assertNotIn(b"thriving", everything)
        self.assertNotIn(b"struggling", everything)
        statuses = {r[[c["name"] for c in t["columns"]].index("version_status")]
                    for t, rows in exp.tables if t["name"] == "approved_values" for r in rows}
        self.assertTrue(statuses <= {"approved", "superseded"}, statuses)

    def test_names_filled_in(self):
        exp = exports.build(self.repo(), "weekly_work", now=NOW)
        rows = read_csv(exports.csv_files(exp)["tasks.csv"])
        h = rows[0]
        first = dict(zip(h, rows[1]))
        self.assertEqual(first["team_name"], "Team A")
        self.assertEqual(first["raised_by_name"], "Jordan Hughes")
        self.assertTrue(first["contributes_to_name"])

    def test_quarter_pack(self):
        repo = self.repo()
        exp = exports.build(repo, "quarter_pack", now=NOW)
        self.assertEqual(exp.quarter_label, "2026-27 Q2", "defaults to the last completed quarter")
        cols = [c["name"] for c in exports.DATASETS["quarter_pack"]["tables"][0]["columns"]]
        self.assertTrue(all(r[cols.index("financial_quarter")] == "2026-27 Q2" for r in exp.table("quarter_values")))
        self.assertTrue(exp.table("quarter_values"))
        weeks = [r[1] for r in exp.table("weekly_updates")]
        self.assertTrue(all(date(2026, 7, 1) <= w <= date(2026, 9, 30) for w in weeks))
        prompt = exp.files["prompt_quarterly_report.md"]
        self.assertIn("quarterly performance report for 2026-27 Q2", prompt)
        self.assertIn("## Without AI", prompt, "works when no AI tool is available")
        self.assertIn("Never invent or estimate numbers", prompt)
        q1 = repo.get("periods", "Q-2026-27-Q1")
        self.assertEqual(exports.build(repo, "quarter_pack", quarter=q1, now=NOW).quarter_label, "2026-27 Q1")

    def test_excel_workbook(self):
        conn = self.conn()
        repo = db.Repo(conn, 1)
        repo.insert("tasks", dict(task_code="TSK-09999", task_name='=HYPERLINK("http://evil.example","Click")',
                                  priority="must", org_unit_key="TM-DATA", raised_by=ADMIN, date_raised=date(2026, 9, 1),
                                  status="open"), ADMIN)
        exp = exports.build(repo, "weekly_work", now=NOW)
        wb = load_workbook(io.BytesIO(exports.to_xlsx(exp)))
        self.assertEqual(wb.sheetnames, ["weekly_updates", "tasks", "problems", "data_dictionary", "readme"])
        ws = wb["tasks"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header[0], "task_code")
        self.assertEqual(ws.freeze_panes, "A2")
        cells = {r[0].value: r for r in ws.iter_rows(min_row=2)}
        evil = cells["TSK-09999"][header.index("task_name")]
        self.assertEqual(evil.data_type, "s", "stored as text, never as a formula")
        raised = cells["TSK-00001"][header.index("date_raised")]
        self.assertIsInstance(raised.value, datetime, "a real date Excel can sort and filter")
        self.assertEqual(raised.number_format, "yyyy-mm-dd")
        dictionary = [[c.value for c in r] for r in wb["data_dictionary"].iter_rows()]
        self.assertEqual(dictionary[0], exports.DICTIONARY_HEADERS)
        self.assertIn(["tasks", "team_name", "text", "Team name."], dictionary)
        readme = "\n".join(str(r[0].value or "") for r in wb["readme"].iter_rows(min_row=2))
        self.assertIn("How to read this data", readme)


class Downloads(AppCase, unittest.TestCase):
    def download(self, **form):
        return self.post("/exports/download", form, csrf_from="/exports/")

    def test_viewer_can_download_excel_and_it_is_logged(self):
        self.login(email=VIEWER)
        r = self.download(dataset="measures", format="xlsx", financial_year="2026-27")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment; filename=\"measures_", r.headers["Content-Disposition"])
        load_workbook(io.BytesIO(r.data))
        log = self.repo().by("export_requests", requested_by=VIEWER)
        self.assertEqual([(x["dataset"], x["filter_label"], x["status"]) for x in log],
                         [("measures", "financial year 2026-27", "done")])
        self.assertIn("Your recent exports", self.client.get("/exports/").get_data(as_text=True))

    def test_zip_has_everything(self):
        self.login(email=STANDARD)
        r = self.download(dataset="quarter_pack", format="both", quarter="Q-2026-27-Q2")
        self.assertEqual(r.mimetype, "application/zip")
        z = zipfile.ZipFile(io.BytesIO(r.data))
        names = {n.split("/", 1)[1] for n in z.namelist()}
        self.assertTrue({"quarter_values.csv", "latest_values.csv", "tasks.csv", "problems.csv", "weekly_updates.csv",
                         "data_dictionary.csv", "README_for_AI.md", "prompt_quarterly_report.md", "export_info.txt",
                         "quarter_pack.xlsx"} <= names, names)
        self.assertEqual(len({n.split("/")[0] for n in z.namelist()}), 1, "one folder inside the zip")

    def test_choices_checked(self):
        self.login(email=STANDARD)
        r = self.download(dataset="audit_log", format="csv")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Choose what to export.", r.get_data(as_text=True))
        r = self.download(dataset="quarter_pack", format="csv", quarter="M-2026-08")
        self.assertIn("Choose a quarter from the list.", r.get_data(as_text=True))
        r = self.download(dataset="measures", format="pdf")
        self.assertIn("Choose a format.", r.get_data(as_text=True))
        self.assertEqual(self.repo().count("export_requests", "requested_by = ?", (STANDARD,)), 0)

    def test_sign_in_needed(self):
        self.assertEqual(self.client.get("/exports/").status_code, 302)
        self.assertEqual(self.client.post("/exports/download").status_code, 400, "CSRF check first")


class Snapshots(AppCase, unittest.TestCase):
    def folder(self):
        return self.tmp / "snapshots"

    def test_due_jobs_write_latest_and_dated_copies(self):
        repo = self.repo()
        ran = exports.run_due_jobs(repo, self.folder(), "default", today=date(2026, 10, 12), now=NOW)
        self.assertEqual(ran, {"EXP-FULL-NIGHTLY": "done", "EXP-WEEKLY-WORK": "done", "EXP-QUARTER-PACK": "done"})
        full = self.folder() / "full_model"
        self.assertTrue((full / "latest" / "values_flat.csv").exists())
        self.assertTrue((full / "latest" / "full_model.xlsx").exists(), "format both")
        self.assertTrue((full / "2026-10-03" / "README_for_AI.md").exists(), "dated copy kept")
        self.assertFalse((self.folder() / "weekly_work" / "latest" / "weekly_work.xlsx").exists(), "csv only")
        self.assertTrue((self.folder() / "quarter_pack" / "latest" / "prompt_quarterly_report.md").exists())
        self.assertEqual(sorted(p.name for p in full.iterdir()), ["2026-10-03", "latest"], "no temporary folders left")
        job = repo.get("export_jobs", "EXP-FULL-NIGHTLY")
        self.assertTrue(job["last_run_status"].startswith("done: 14 tables"))
        self.assertEqual(exports.run_due_jobs(repo, self.folder(), "default", today=date(2026, 10, 3), now=NOW), {},
                         "nothing due again the same day")

    def test_rerun_replaces_latest(self):
        repo = self.repo()
        job = repo.get("export_jobs", "EXP-WEEKLY-WORK")
        exports.run_job(repo, job, self.folder(), "default", now=NOW)
        stale = self.folder() / "weekly_work" / "latest" / "stale.txt"
        stale.write_text("old")
        exports.run_job(repo, job, self.folder(), "default", now=NOW)
        self.assertFalse(stale.exists())

    def test_failure_recorded_and_others_still_run(self):
        conn = self.conn()
        conn.execute("UPDATE export_jobs SET folder_path = '../escape' WHERE job_code = 'EXP-WEEKLY-WORK'")
        repo = db.Repo(conn, 1)
        ran = exports.run_due_jobs(repo, self.folder(), "default", today=date(2026, 10, 12), now=NOW)
        self.assertEqual(ran["EXP-WEEKLY-WORK"], "failed")
        self.assertEqual(ran["EXP-FULL-NIGHTLY"], "done")
        self.assertFalse((self.tmp / "escape").exists())
        self.assertIn("folder name", repo.get("export_jobs", "EXP-WEEKLY-WORK")["last_run_status"])

    def test_run_now_admin_only(self):
        self.login(email=STANDARD)
        self.assertEqual(self.client.get("/exports/scheduled").status_code, 403)
        self.client = self.app.test_client()
        self.login()
        r = self.post("/exports/scheduled/run", dict(job_code="EXP-WEEKLY-WORK"), csrf_from="/exports/scheduled")
        self.assertEqual(r.status_code, 302)
        self.assertTrue((self.folder() / "weekly_work" / "latest" / "tasks.csv").exists())

    def test_first_request_of_the_day_runs_them(self):
        self.app.config["DISABLE_DAILY_JOBS"] = False
        self.login(email=STANDARD)
        self.client.get("/")
        self.assertTrue((self.folder() / "full_model" / "latest" / "values_flat.csv").exists())

    def test_admin_form_checks_jobs(self):
        self.login()
        form = dict(job_code="EXP-NEW", job_name="New", dataset="measures", format="csv", frequency="weekly",
                    run_day="9", folder_path="a/b", keep_history="1", active="1", last_run_status="hacked")
        r = self.post("/admin/export_jobs/new", form)
        html = r.get_data(as_text=True)
        self.assertIn("Enter a day of the week from 1 (Monday) to 7 (Sunday).", html)
        self.assertIn("Use only letters, numbers, - and _", html)
        self.post("/admin/export_jobs/new", dict(form, run_day="1", folder_path="measures_weekly"))
        job = self.repo().get("export_jobs", "EXP-NEW")
        self.assertEqual(job["folder_path"], "measures_weekly")
        self.assertIsNone(job["last_run_status"], "written only by the app")


class DataLinks(AppCase, unittest.TestCase):
    def make(self, dataset="measures", name="Board dashboard"):
        self.login()
        r = self.post("/exports/links", dict(name=name, dataset=dataset), csrf_from="/exports/links")
        m = re.search(r'value="(http://pms\.test/data/[^"]+/)"', r.get_data(as_text=True))
        self.client = self.app.test_client()  # BI tools don't sign in
        return m.group(1).replace("http://pms.test", "")

    def test_live_csv_without_signing_in(self):
        url = self.make()
        r = self.client.get(url + "values_flat.csv")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.mimetype, "text/csv")
        self.assertEqual(r.headers["Cache-Control"], "no-store")
        rows = read_csv(r.data)
        self.assertEqual(len(rows) - 1, self.repo().count("rpt_values"))
        dictionary = read_csv(self.client.get(url + "data_dictionary.csv").data)
        self.assertIn("financial_quarter", [r[1] for r in dictionary if r[0] == "values_flat"])
        self.assertIn(b"How to read this data", self.client.get(url + "README_for_AI.md").data)
        self.assertIn("values_flat.csv", self.client.get(url).get_data(as_text=True))
        row = self.conn().execute("SELECT use_count, last_used_at FROM data_links").fetchone()
        self.assertEqual(row[0], 4)
        self.assertIsNotNone(row[1])

    def test_only_its_own_files(self):
        url = self.make()
        self.assertEqual(self.client.get(url + "tasks.csv").status_code, 404)
        self.assertEqual(self.client.get(url + "wellbeing_checkins.csv").status_code, 404)
        self.assertEqual(self.client.get(url + "..%2Fsecret_key").status_code, 404)
        self.assertEqual(self.client.get("/data/not-a-real-token/values_flat.csv").status_code, 404)

    def test_token_stored_hashed_and_shown_once(self):
        url = self.make()
        token = url.split("/")[2]
        stored = self.conn().execute("SELECT token_hash FROM data_links").fetchone()[0]
        self.assertNotIn(token, stored)
        self.login()
        self.assertNotIn(token, self.client.get("/exports/links").get_data(as_text=True))
        audit = self.repo().by("audit_log", list_name="data_links")
        self.assertEqual([a["action"] for a in audit], ["create"])

    def test_turning_off(self):
        url = self.make()
        self.login()
        link_id = self.conn().execute("SELECT id FROM data_links").fetchone()[0]
        self.post("/exports/links/revoke", dict(id=link_id), csrf_from="/exports/links")
        self.assertEqual(self.client.get(url + "values_flat.csv").status_code, 404)
        self.assertEqual(len(self.repo().by("audit_log", list_name="data_links")), 2)

    def test_admin_only(self):
        self.login(email=STANDARD)
        self.assertEqual(self.client.get("/exports/links").status_code, 403)
        r = self.post("/exports/links", dict(name="x", dataset="measures"), csrf_from="/exports/")
        self.assertEqual(r.status_code, 403)
        self.client = self.app.test_client()
        self.login()
        r = self.post("/exports/links", dict(name="", dataset="quarter_pack"), csrf_from="/exports/links")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Enter a name", r.get_data(as_text=True))
        self.assertEqual(self.conn().execute("SELECT COUNT(*) FROM data_links").fetchone()[0], 0)


if __name__ == "__main__":
    unittest.main()
